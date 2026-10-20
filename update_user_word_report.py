# -*- coding: utf-8 -*-
"""
Kullanıcının düzenlediği Word raporunu (IHA_Telemetri_Teknik_Rapor_USER_BACKUP.docx)
hiçbir metin değişikliği kaybolmadan yerinde (in-place) günceller:
1. En başa Selçuk Üniversitesi Teknoloji Fakültesi, AAES ve Drohytech logolu resmi Kapak Sayfası ekler.
2. Kapak sayfasında tam başlık ve künye tablosu yer aldığı için 2. sayfa başındaki mükerrer (tekrar eden)
   eski başlık ve künye tablosunu kaldırarak 2. sayfanın doğrudan ÖZET (ABSTRACT) ile başlamasını sağlar.
3. Rapor içindeki 10 adet sarı ekran görüntüsü yer tutucu kutusunu gerçek ekran görüntüleriyle değiştirir.
4. Tablo 4.2'deki dosya/kayıt sayılarını gömülen gerçek ekran görüntüleriyle (telemetry_mixed.csv: 30 geçerli, 10 hatalı) birebir uyumlu hale getirir.
5. Tüm tablo satırlarına cantSplit (satırın sayfa ortasında bölünmesini engelleme), veri tablolarının ilk satırına tblHeader
   (sayfa taşmasında başlık tekrarı) ve tüm alt başlıklara/görsellere keep_with_next uygular.
6. 9.5 pt font boyutlarını 9 pt, 10.5 pt font boyutlarını 10 pt olarak standartlaştırır.
"""
import os
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
BACKUP_DOCX = os.path.join(PROJECT_DIR, "IHA_Telemetri_Teknik_Rapor_USER_BACKUP.docx")
OUT_DOCX_MAIN = os.path.join(PROJECT_DIR, "IHA_Telemetri_Teknik_Rapor_Mehmet_Akif_Urey.docx")

PLACEHOLDER_MAP = {
    "Şekil 1.1": "01_csv_senaryo_uretici_arayuzu.png",
    "Şekil 2.1": "02_ham_csv_orijinal_sekmesi.png",
    "Şekil 3.1": "03_gecerli_ucus_analiz_ve_grafikler.png",
    "Şekil 4.1": "04_hata_dosya_secilmeden_analiz.png",
    "Şekil 4.2": "05_hata_bos_veya_bozuk_dosya.png",
    "Şekil 4.3": "06_hata_eksik_sutunlu_dosya.png",
    "Şekil 4.4": "07_kismi_hatali_uyari_popup.png",
    "Şekil 4.5": "08_hatali_kayitlar_raporu_sekmesi.png",
    "Şekil 5.1": "09_zaman_boslugu_gap_grafik_anomalisi.png",
    "Şekil 6.1": "10_pytest_35_otomatik_test_ciktisi.png",
}


def set_cell_shading(cell, color_hex: str):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith("shd"):
            tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=80, bottom=80, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith("tcMar"):
            tcPr.remove(child)
    tcMar = OxmlElement("w:tcMar")
    for m, val in [("top", top), ("bottom", bottom), ("left", left), ("right", right)]:
        node = OxmlElement(f"w:{m}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        tcMar.append(node)
    tcPr.append(tcMar)


def set_table_borders(table, color="CBD5E1", sz="6", val="single"):
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith("tblBorders"):
            tblPr.remove(child)
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f"</w:tblBorders>"
    )
    tblPr.append(borders)


def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    for child in list(trPr):
        if child.tag.endswith("cantSplit"):
            trPr.remove(child)
    trPr.append(OxmlElement("w:cantSplit"))


def set_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    for child in list(trPr):
        if child.tag.endswith("tblHeader"):
            trPr.remove(child)
    trPr.append(OxmlElement("w:tblHeader"))


def build_cover_page_elements(doc):
    """
    Kapak sayfası elemanlarını geçici olarak oluşturur ve XML düğümlerini (list) döndürür;
    ardından bu düğümler dokümanın en başına yerleştirilir.
    """
    created_nodes = []

    def new_p(align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        created_nodes.append(p._p)
        return p

    # 1. Üst Logolar Tablosu (Selçuk Üniversitesi | AAES | Drohytech)
    logo_tbl = doc.add_table(rows=1, cols=3)
    logo_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(logo_tbl, color="FFFFFF", sz="0", val="none")
    created_nodes.append(logo_tbl._tbl)

    logos_info = [
        (
            os.path.join(PROJECT_DIR, "assets", "logos", "cover_selcuk.png"),
            "T.C. SELÇUK ÜNİVERSİTESİ\nTeknoloji Fakültesi",
        ),
        (
            os.path.join(PROJECT_DIR, "assets", "logos", "cover_aaes.png"),
            "AAES İLERİ MÜHENDİSLİK\nSavunma ve Havacılık San. Tic. Ltd. Şti.",
        ),
        (
            os.path.join(PROJECT_DIR, "assets", "logos", "cover_drohytech.png"),
            "DROHYTECH\nİHA ve Otonom Sistem Teknolojileri",
        ),
    ]

    for idx, (img_path, caption_text) in enumerate(logos_info):
        cell = logo_tbl.rows[0].cells[idx]
        cell.width = Inches(2.15)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        set_cell_shading(cell, "FFFFFF")
        set_cell_margins(cell, top=30, bottom=30, left=40, right=40)

        p_img = cell.paragraphs[0]
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(0)
        p_img.paragraph_format.space_after = Pt(3)
        if os.path.exists(img_path):
            r_img = p_img.add_run()
            r_img.add_picture(img_path, width=Inches(1.88))

        p_cap = cell.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_before = Pt(0)
        p_cap.paragraph_format.space_after = Pt(0)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(8)
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(30, 58, 138)

    # 2. Kurumsal Üst Başlık Bloğu
    p_tc = new_p(space_before=8, space_after=2)
    r = p_tc.add_run("T.C.\nSELÇUK ÜNİVERSİTESİ\nTEKNOLOJİ FAKÜLTESİ\nBİLGİSAYAR MÜHENDİSLİĞİ BÖLÜMÜ")
    r.font.name = "Calibri"
    r.font.size = Pt(13)
    r.font.bold = True
    r.font.color.rgb = RGBColor(15, 44, 89)

    p_sub_hdr = new_p(space_before=3, space_after=10)
    r2 = p_sub_hdr.add_run("İŞLETMEDE MESLEKİ EĞİTİM (İME) — TEKNİK PROJE VE ARAŞTIRMA RAPORU")
    r2.font.name = "Calibri"
    r2.font.size = Pt(10)
    r2.font.bold = True
    r2.font.color.rgb = RGBColor(37, 99, 235)

    # 3. Ana Rapor Başlığı Kutusu (Şık Çerçeveli Kart)
    title_tbl = doc.add_table(rows=1, cols=1)
    title_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(title_tbl, color="1E3A8A", sz="12", val="single")
    created_nodes.append(title_tbl._tbl)

    t_cell = title_tbl.rows[0].cells[0]
    t_cell.width = Inches(6.45)
    set_cell_shading(t_cell, "F8FAFC")
    set_cell_margins(t_cell, top=150, bottom=150, left=180, right=180)

    tp1 = t_cell.paragraphs[0]
    tp1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp1.paragraph_format.space_before = Pt(2)
    tp1.paragraph_format.space_after = Pt(6)
    tr1 = tp1.add_run(
        "İHA TELEMETRİ ANALİZ VE GÖRSELLEŞTİRME UYGULAMASI:\n"
        "LİTERATÜR TARAMASI, SİSTEM MİMARİSİ VE TEKNİK TASARIM RAPORU"
    )
    tr1.font.name = "Calibri"
    tr1.font.size = Pt(15)
    tr1.font.bold = True
    tr1.font.color.rgb = RGBColor(15, 44, 89)

    tp2 = t_cell.add_paragraph()
    tp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp2.paragraph_format.space_before = Pt(2)
    tp2.paragraph_format.space_after = Pt(6)
    tr2 = tp2.add_run(
        "Uçuş Sonrası (Post-Flight) Veri Doğrulama, Kısmi Kurtarma (Partial Recovery),\n"
        "Zaman Serisi Anomali Analizi ve Bağımsız CSV Senaryo Üretici Ekosistemi"
    )
    tr2.font.name = "Calibri"
    tr2.font.size = Pt(11)
    tr2.font.italic = True
    tr2.font.color.rgb = RGBColor(51, 65, 85)

    tp3 = t_cell.add_paragraph()
    tp3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp3.paragraph_format.space_before = Pt(2)
    tp3.paragraph_format.space_after = Pt(2)
    tr3 = tp3.add_run("Sürüm v1.6.0  •  Otomatik Doğrulama: 35 Pytest Senaryosu (%100 Başarılı)")
    tr3.font.name = "Calibri"
    tr3.font.size = Pt(9)
    tr3.font.bold = True
    tr3.font.color.rgb = RGBColor(21, 128, 61)

    # Boşluk
    new_p(space_before=8, space_after=4)

    # 4. Künye ve Kurumsal Teslimat Tablosu (Eski Tablo 0'daki tüm bilgileri de kapsar)
    meta_rows = [
        ("Hazırlayan Öğrenci", "Mehmet Akif ÜREY (Öğrenci No: 233311066)"),
        ("Üniversite / Fakülte / Bölüm", "Selçuk Üniversitesi / Teknoloji Fakültesi / Bilgisayar Mühendisliği Bölümü"),
        ("İlgili Kurum / Şirket", "AAES İleri Mühendislik San. ve Tic. Ltd. Şti. — Drohytech İHA Teknolojileri"),
        ("Değerlendirici / Proje Yürütücüsü", "Can YILMAZ"),
        ("Teslim Tarihi", "20 Ekim 2026 Salı"),
        ("Yazılım Sürümü ve Teknoloji Yığını", "v1.6.0 — Python 3.12, CustomTkinter, Pandas, NumPy, Matplotlib, Pillow, Pytest"),
        ("Ana Proje ve Ek Uygulama Yapısı", "İHA Telemetri Analiz İstasyonu (Ana Dizin) & CSV Senaryo Üretici (csv_senaryo_uretici/)"),
        ("Otomatik Test ve Görsel Kapsamı", "35 Otomatik Test (10 Zorunlu + 5 Ek Senaryo + 20 Birim Testi) + 10 Canlı Ekran Görüntüsü"),
    ]

    meta_tbl = doc.add_table(rows=len(meta_rows) + 1, cols=2)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(meta_tbl, color="CBD5E1", sz="6", val="single")
    created_nodes.append(meta_tbl._tbl)

    hdr_cells = meta_tbl.rows[0].cells
    hdr_cells[0].width = Inches(2.35)
    hdr_cells[1].width = Inches(4.10)
    for ci, h_txt in enumerate(["Kurumsal Künye ve Proje Bilgisi", "Detay"]):
        set_cell_shading(hdr_cells[ci], "0F2C59")
        set_cell_margins(hdr_cells[ci], top=60, bottom=60, left=100, right=100)
        hp = hdr_cells[ci].paragraphs[0]
        hp.paragraph_format.space_before = Pt(1)
        hp.paragraph_format.space_after = Pt(1)
        hr = hp.add_run(h_txt)
        hr.font.name = "Calibri"
        hr.font.size = Pt(9)
        hr.font.bold = True
        hr.font.color.rgb = RGBColor(255, 255, 255)

    for ri, (k, v) in enumerate(meta_rows, start=1):
        rcells = meta_tbl.rows[ri].cells
        rcells[0].width = Inches(2.35)
        rcells[1].width = Inches(4.10)
        bg = "F8FAFC" if ri % 2 == 1 else "FFFFFF"
        for ci, val_txt in enumerate([k, v]):
            set_cell_shading(rcells[ci], bg)
            set_cell_margins(rcells[ci], top=55, bottom=55, left=100, right=100)
            cp = rcells[ci].paragraphs[0]
            cp.paragraph_format.space_before = Pt(1)
            cp.paragraph_format.space_after = Pt(1)
            cr = cp.add_run(val_txt)
            cr.font.name = "Calibri"
            cr.font.size = Pt(9)
            cr.font.bold = (ci == 0)
            cr.font.color.rgb = RGBColor(15, 44, 89) if ci == 0 else RGBColor(30, 41, 59)

    # 5. Alt Tarih/Şehir ve Sayfa Sonu (Page Break)
    p_footer = new_p(space_before=12, space_after=0)
    rf = p_footer.add_run("KONYA — EKİM 2026")
    rf.font.name = "Calibri"
    rf.font.size = Pt(11)
    rf.font.bold = True
    rf.font.color.rgb = RGBColor(15, 44, 89)

    p_brk = doc.add_paragraph()
    p_brk.paragraph_format.space_before = Pt(0)
    p_brk.paragraph_format.space_after = Pt(0)
    r_brk = p_brk.add_run()
    r_brk.add_break(docx.enum.text.WD_BREAK.PAGE)
    created_nodes.append(p_brk._p)

    return created_nodes


def update_cell_text_preserving_style(cell, new_text):
    p = cell.paragraphs[0]
    if p.runs:
        first_run = p.runs[0]
        fname = first_run.font.name or "Calibri"
        fbold = first_run.font.bold
        fcolor = first_run.font.color.rgb if first_run.font.color else None
        for r in list(p.runs):
            p._p.remove(r._r)
        nr = p.add_run(new_text)
        nr.font.name = fname
        nr.font.size = Pt(9)
        nr.font.bold = fbold
        if fcolor:
            nr.font.color.rgb = fcolor
    else:
        cell.text = new_text


def main():
    print(f"Yedek dosya yükleniyor: {BACKUP_DOCX}")
    doc = docx.Document(BACKUP_DOCX)

    # 1. Kapak eklendiği için 2. sayfanın en başındaki mükerrer başlık + (Sürüm v1.6.0) + eski Proje Bilgisi tablosunu kaldır
    # (Body[0]: Başlık, Body[1]: (Sürüm v1.6.0), Body[2]: Tablo 0 Proje Bilgisi, Body[3]: boş paragraf)
    body_children = list(doc.element.body)
    for el in body_children[:4]:
        txt = "".join(el.itertext()).strip()
        if "ÖZET" not in txt:
            doc.element.body.remove(el)
    print("[MÜKERRER BAŞLIK TEMİZLENDİ] 2. sayfa doğrudan ÖZET (ABSTRACT) ile başlayacak şekilde düzenlendi.")

    # 2. Tablo 4.2'deki dosya ve kayıt sayılarını gömülen gerçek ekran görüntüleriyle birebir eşleştir
    for tbl in doc.tables:
        if len(tbl.rows) >= 5 and len(tbl.rows[0].cells) == 4:
            hdr0 = tbl.rows[0].cells[0].text.strip()
            if hdr0 == "Senaryo":
                # Senaryo 2
                update_cell_text_preserving_style(
                    tbl.rows[2].cells[1],
                    "'telemetry_empty.csv' (veya '13_bos_dosya_0_bayt.csv') seçilip analiz edilmesi.",
                )
                update_cell_text_preserving_style(
                    tbl.rows[2].cells[3],
                    "Doğrulama / Yükleme Hatası (Kritik): 'Dosya işlenemedi! Neden: CSV dosyası boş (0 bayt).'",
                )
                # Senaryo 3
                update_cell_text_preserving_style(
                    tbl.rows[3].cells[1],
                    "'telemetry_missing_columns.csv' (veya '11_eksik_sutunlu_csv.csv') yüklenip analiz edilmesi.",
                )
                update_cell_text_preserving_style(
                    tbl.rows[3].cells[3],
                    "Doğrulama / Yükleme Hatası (Kritik): 'Dosya işlenemedi! Neden: Gerekli sütunlar eksik: battery_pct, pitch, yaw'",
                )
                # Senaryo 4
                update_cell_text_preserving_style(
                    tbl.rows[4].cells[1],
                    "'telemetry_mixed.csv' (40 toplam: 30 geçerli, 10 hatalı satır) analiz edilmesi.",
                )
                update_cell_text_preserving_style(
                    tbl.rows[4].cells[2],
                    "30 geçerli satır grafik ve tabloya çizilir; 10 hatalı satır hata sekmesine aktarılır.",
                )
                update_cell_text_preserving_style(
                    tbl.rows[4].cells[3],
                    "Doğrulama Uyarısı: 'Telemetri dosyasında 10 adet hatalı satır tespit edildi ve elendi. İşlenen Geçerli Kayıt: 30, Reddedilen Hatalı Kayıt: 10'",
                )
                print("[TABLO 4.2 UYUMLANDIRILDI] Senaryo 2, 3 ve 4 metinleri ekran görüntüleriyle birebir eşleştirildi.")

    # 3. Yer tutucu tabloları gerçek ekran görüntüleriyle yerinde (in-place) değiştir
    replaced_count = 0
    for tbl in doc.tables:
        cell_text = tbl.rows[0].cells[0].text.strip()
        if "EKRAN GÖRÜNTÜSÜ" in cell_text or "YER TUTUCU" in cell_text or "📸" in cell_text:
            matched_fig = None
            for fig_key, png_name in PLACEHOLDER_MAP.items():
                if fig_key in cell_text:
                    matched_fig = (fig_key, png_name)
                    break
            if matched_fig:
                fig_key, png_name = matched_fig
                png_path = os.path.join(PROJECT_DIR, "screenshots", png_name)
                if os.path.exists(png_path):
                    cell = tbl.rows[0].cells[0]
                    cell.text = ""
                    set_cell_shading(cell, "FFFFFF")
                    set_cell_margins(cell, top=50, bottom=50, left=50, right=50)
                    set_table_borders(tbl, color="CBD5E1", sz="6", val="single")
                    set_row_cant_split(tbl.rows[0])
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    p.paragraph_format.keep_with_next = True
                    run = p.add_run()
                    run.add_picture(png_path, width=Inches(5.95))
                    replaced_count += 1
                    print(f"[GÖRSEL EKLENDİ] {fig_key} -> {png_name}")

    print(f"Toplam yerleştirilen ekran görüntüsü: {replaced_count}/10")

    # 4. Tüm tablolarda satır bölünmesini (cantSplit) engelle, başlık satırlarını (tblHeader) işaretle,
    # kısa kontrol listesi tablolarının sayfa ortasında bölünmesini (keep_with_next) önle
    for tbl in doc.tables:
        num_rows = len(tbl.rows)
        if num_rows > 1:
            set_row_header(tbl.rows[0])
        for ri, row in enumerate(tbl.rows):
            set_row_cant_split(row)
            # Kısa tabloların (örn. 4-5 satırlık CHK kontrol tabloları) sayfa sonunda bölünmemesi için
            if num_rows <= 6 and ri < num_rows - 1:
                for cell in row.cells:
                    for cp in cell.paragraphs:
                        cp.paragraph_format.keep_with_next = True

    # 5. Tüm ana ve alt başlıkların (1., 1.A., 1.B., 1.C. vb.) sonraki paragraf/tablo ile aynı sayfada kalmasını sağla
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        is_heading = (
            (p.style and "Heading" in p.style.name)
            or txt.startswith(("1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9."))
            or txt in ("ÖZET (ABSTRACT)", "PROJE GENEL GEREKSİNİM KONTROL LİSTESİ")
            or txt.startswith("Karşılaşılan Sorun")
            or txt.startswith("İzlenen Çözüm Yolu")
        )
        if is_heading:
            p.paragraph_format.keep_with_next = True

    # 6. Kapak sayfasını oluştur ve dokümanın en başına taşı
    first_body_el = doc.element.body[0]
    cover_nodes = build_cover_page_elements(doc)
    for node in cover_nodes:
        first_body_el.addprevious(node)
    print("[KAPAK EKLENDİ] Selçuk Üniversitesi + AAES + Drohytech logolu kapak sayfası en başa eklendi.")

    # 7. Font boyutlarını güncelle: 9.5 pt -> 9 pt, 10.5 pt -> 10 pt
    if "Normal" in doc.styles:
        doc.styles["Normal"].font.size = Pt(10)

    p_95_to_9 = 0
    p_105_to_10 = 0

    for p in doc.paragraphs:
        for r in p.runs:
            if r.font and r.font.size:
                pt_val = round(r.font.size.pt, 1)
                if pt_val == 9.5:
                    r.font.size = Pt(9)
                    p_95_to_9 += 1
                elif pt_val == 10.5:
                    r.font.size = Pt(10)
                    p_105_to_10 += 1
                elif pt_val == 12.5:
                    r.font.size = Pt(12)

    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.font and r.font.size:
                            pt_val = round(r.font.size.pt, 1)
                            if pt_val == 9.5:
                                r.font.size = Pt(9)
                                p_95_to_9 += 1
                            elif pt_val == 10.5:
                                r.font.size = Pt(10)
                                p_105_to_10 += 1

    print(f"[FONT GÜNCELLENDİ] 9.5->9 pt ({p_95_to_9} run), 10.5->10 pt ({p_105_to_10} run)")

    # 8. Ana proje dizinine kaydet ve diğer klasörlere kopyala
    doc.save(OUT_DOCX_MAIN)
    print(f"[KAYDEDİLDİ] {OUT_DOCX_MAIN} ({os.path.getsize(OUT_DOCX_MAIN):,} bayt)")

    targets = [
        r"C:\Users\PC\Desktop\IHA_Masaustu_Uygulamalari\IHA_Telemetri_Teknik_Rapor_Mehmet_Akif_Urey.docx",
        r"C:\Users\PC\Desktop\IHA_Proje_Teslim_Paketi_Mehmet_Akif_Urey\IHA_Telemetri_Teknik_Rapor_Mehmet_Akif_Urey.docx",
    ]
    for t in targets:
        try:
            shutil.copy2(OUT_DOCX_MAIN, t)
            print(f"[SENKRONİZE EDİLDİ] {t}")
        except PermissionError:
            alt_t = t.replace(".docx", "_GUNCEL_KAPAKLI_VE_GORSELLI.docx")
            shutil.copy2(OUT_DOCX_MAIN, alt_t)
            print(f"[AÇIK DOSYA KİLİDİ] Word'de açık olduğu için ayrıca kaydedildi: {alt_t}")


if __name__ == "__main__":
    main()
